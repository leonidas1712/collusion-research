"""Apply the Apart template's direct formatting and native front matter."""
from copy import deepcopy
from zipfile import ZipFile
from lxml import etree
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, Inches
from docx.text.paragraph import Paragraph


def copy_properties(paragraph, reference):
    old = paragraph._p.pPr
    if old is not None:
        paragraph._p.remove(old)
    if reference._p.pPr is not None:
        paragraph._p.insert(0, deepcopy(reference._p.pPr))


def replace_text(paragraph, text, caps=False):
    sample = next((r for r in paragraph.runs if r.text), paragraph.runs[0])
    rpr = deepcopy(sample._r.rPr)
    for child in list(paragraph._p):
        if child.tag != qn('w:pPr'):
            paragraph._p.remove(child)
    run = paragraph.add_run(text)
    if rpr is not None:
        run._r.insert(0, rpr)
    if caps:
        run.font.all_caps = True
    return run


def apply_layout(doc, template):
    reference = Document(template)
    paragraphs = doc.paragraphs
    title = paragraphs[0].text
    abstract_at = next(i for i,p in enumerate(paragraphs) if p.text == 'Abstract')
    intro_at = next(i for i,p in enumerate(paragraphs) if p.text.startswith('1. Introduction'))
    author = '\n'.join(p.text for p in paragraphs[1:abstract_at] if p.text)
    if ', ' in author:
        author = author.replace(', ', '\n', 1)
    abstract = paragraphs[abstract_at+1:intro_at]
    main_heading = next(p for p in reference.paragraphs if p.text == '1. Introduction')
    subheading = next(p for p in reference.paragraphs if p.text == 'Limitations')
    for p in paragraphs[intro_at:]:
        if p.style.name.startswith('Heading'):
            sample = main_heading if p.style.name == 'Heading 2' else subheading
            copy_properties(p, sample)
            p.paragraph_format.keep_with_next = True
            for r in p.runs:
                r.font.size = Pt(14 if sample == main_heading else 13)
                r.bold = True
        if p.text in ('References', 'Appendix'):
            p.paragraph_format.page_break_before = True
    front = deepcopy(reference.tables[0]._tbl)
    front_paragraphs = [Paragraph(p, doc) for p in front.iter(qn('w:p'))]
    title_p = next(p for p in front_paragraphs if p.text == 'PROJECT TITLE')
    footnote = deepcopy(title_p._p.find('.//' + qn('w:footnoteReference')))
    title_run = replace_text(title_p, title)
    title_run.font.all_caps = False
    r = title_p.add_run()
    r._r.append(footnote)
    author_p = next(p for p in front_paragraphs if p.text.startswith('Author name 1'))
    author_node = deepcopy(author_p._p)
    replace_text(Paragraph(author_node, doc), author)
    nested = list(front.iter(qn('w:tbl')))[1]
    with_node = deepcopy(next(p._p for p in front_paragraphs if p.text == 'With\nApart Research'))
    nested.getparent().replace(nested, author_node)
    author_node.addnext(with_node)
    abstract_p = next(p for p in front_paragraphs if p.text.startswith('Summarize your project'))
    cell = abstract_p._p.getparent()
    offset = cell.index(abstract_p._p)
    for p in abstract:
        node = deepcopy(p._p)
        clone = Paragraph(node, doc)
        copy_properties(clone, abstract_p)
        for r in clone.runs:
            r.italic = True
            r.font.size = Pt(11)
        cell.insert(offset, node)
        offset += 1
    cell.remove(abstract_p._p)
    for p in paragraphs[:intro_at]:
        p._p.getparent().remove(p._p)
    body = doc.element.body
    body.insert(0, deepcopy(reference.paragraphs[0]._p))
    body.insert(1, front)
    body.insert(2, deepcopy(reference.paragraphs[1]._p))
    # Citeproc supplies alphabetical author-year entries; keep them unnumbered.
    numbering = doc.part.numbering_part.element
    for p in doc.paragraphs:
        if p.style.name == 'Bibliography':
            for num in p._p.xpath('./w:pPr/w:numPr'):
                num.getparent().remove(num)
            p.paragraph_format.left_indent = Inches(.25)
            p.paragraph_format.first_line_indent = Inches(-.25)
            p.paragraph_format.space_after = Pt(4)
    # Word permits block-level bookmarks; LibreOffice can misplace their
    # PDF destinations. Anchor each reference inside its own paragraph.
    body = doc.element.body
    for mark in list(body.findall(qn('w:bookmarkStart'))):
        following = mark.getnext()
        if following is None or following.tag != qn('w:p'):
            continue
        end = following.getnext()
        if end is None or end.tag != qn('w:bookmarkEnd') or end.get(qn('w:id')) != mark.get(qn('w:id')):
            continue
        following.insert(1 if following.pPr is not None else 0, mark)
        following.append(end)
    # Match the template's dash bullets without introducing Symbol font.
    for level in numbering.iter(qn('w:lvl')):
        fonts = list(level.iter(qn('w:rFonts')))
        if any(f.get(qn('w:ascii')) == 'Symbol' for f in fonts):
            level.find(qn('w:lvlText')).set(qn('w:val'), '-')
            for f in fonts:
                for attr in ['ascii','hAnsi','cs','eastAsia']:
                    f.set(qn('w:'+attr), 'Old Standard TT')
    # Old Standard TT lacks kappa in this installed font; use the template's
    # Arial explicitly for that symbol instead of a platform-selected fallback.
    for run in list(doc.element.body.iter(qn('w:r'))):
        texts = list(run.iter(qn('w:t')))
        if len(texts) != 1 or 'κ' not in (texts[0].text or ''):
            continue
        text = texts[0].text
        parent = run.getparent(); at = parent.index(run)
        pieces = text.split('κ')
        for i, piece in enumerate(pieces):
            if i:
                node = deepcopy(run); node.find(qn('w:t')).text = 'κ'
                fonts = node.get_or_add_rPr().get_or_add_rFonts()
                for attr in ['ascii','hAnsi','cs','eastAsia']:
                    fonts.set(qn('w:'+attr), 'Arial')
                parent.insert(at,node); at += 1
            if piece:
                node = deepcopy(run); node.find(qn('w:t')).text = piece
                parent.insert(at,node); at += 1
        parent.remove(run)
    return footnote.get(qn('w:id'))


def copy_title_footnote(output, template, old_id):
    with ZipFile(output) as z:
        parts = {n:z.read(n) for n in z.namelist()}
    with ZipFile(template) as z:
        source = etree.fromstring(z.read('word/footnotes.xml'))
        source_rels = etree.fromstring(z.read('word/_rels/footnotes.xml.rels'))
    notes = etree.fromstring(parts['word/footnotes.xml'])
    new_id = max(int(n.get(qn('w:id'))) for n in notes) + 1
    note = deepcopy(next(n for n in source if n.get(qn('w:id')) == old_id))
    note.set(qn('w:id'), str(new_id))
    rel_ns = 'http://schemas.openxmlformats.org/package/2006/relationships'
    rels = etree.fromstring(parts['word/_rels/footnotes.xml.rels']) if 'word/_rels/footnotes.xml.rels' in parts else etree.Element('{'+rel_ns+'}Relationships', nsmap={None:rel_ns})
    for rel in source_rels:
        old = rel.get('Id'); new = 'rIdTemplate' + old
        rel.set('Id', new); rels.append(rel)
        for link in note.iter(qn('w:hyperlink')):
            if link.get(qn('r:id')) == old:
                link.set(qn('r:id'), new)
    notes.append(note)
    root = etree.fromstring(parts['word/document.xml'])
    for ref in root.iter(qn('w:footnoteReference')):
        if ref.get(qn('w:id')) == old_id:
            ref.set(qn('w:id'), str(new_id))
    for name,element in [('word/document.xml',root),('word/footnotes.xml',notes),('word/_rels/footnotes.xml.rels',rels)]:
        parts[name] = etree.tostring(element, xml_declaration=True, encoding='UTF-8', standalone=True)
    with ZipFile(output,'w',compression=8) as z:
        for name,data in parts.items():
            z.writestr(name,data)
