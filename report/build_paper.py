#!/usr/bin/env python3
"""Render paper.md using Pandoc citeproc and the Apart reference document."""
from __future__ import annotations
import argparse
import hashlib
import json
import re
import shutil
import subprocess
import sys
import tempfile
from xml.etree import ElementTree as ET
from pathlib import Path
from zipfile import ZipFile
from lxml import etree
from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from pypdf import PdfReader
from template_layout import apply_layout, copy_title_footnote
from build_docx import (ROOT, set_cell_margins, set_cell_shading, set_table_borders,
                        set_repeat_header, set_run_font)

RENDERER = None


def digest(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()


def walk(node):
    if isinstance(node,dict):
        yield node
        for value in node.values():yield from walk(value)
    elif isinstance(node,list):
        for value in node:yield from walk(value)


def prepare_ast(source,raw):
    ast=json.loads(subprocess.check_output(['pandoc','-f','markdown-implicit_figures','-t','json'],input=raw.decode(),text=True))
    blocks=ast['blocks'];start=next(i for i,b in enumerate(blocks) if b['t']=='Header' and b['c'][0]==1)
    blocks=blocks[start:]
    # The source's thematic separators are editorial section separators, not page breaks.
    blocks=[b for b in blocks if b['t']!='HorizontalRule']
    for i,b in enumerate(blocks):
        if b['t']=='Header' and any(n.get('c')=='References' for n in walk(b) if n.get('t')=='Str'):
            blocks.insert(i+1,{'t':'Div','c':[['refs',[],[]],[]]});break
    if not any(b.get('t')=='Div' and b['c'][0][0]=='refs' for b in blocks):
        at=next((i for i,b in enumerate(blocks) if b['t']=='Header' and b['c'][1][0]=='appendix'),len(blocks))
        blocks[at:at]=[{'t':'Header','c':[2,['references',[],[]],[{'t':'Str','c':'References'}]]},{'t':'Div','c':[['refs',[],[]],[]]}]
    ast['blocks']=blocks
    images=[]
    for b in walk(ast):
        if b.get('t')=='Image':
            path=source.parent/b['c'][2][0]
            if not path.is_file():raise FileNotFoundError(path)
            from PIL import Image
            with Image.open(path) as im:w,h=im.size
            width=min(6.5,6.3*w/h)
            b['c'][0][2].append(['width',f'{width}in'])
            images.append({'path':str(path.relative_to(ROOT)),'sha256':digest(path)})
    return ast,images


def format_docx(path, template):
    # Pandoc copies embedded TTFs from this template but omits their content-type default.
    with ZipFile(path) as z:
        parts={n:z.read(n) for n in z.namelist()}
    ns='http://schemas.openxmlformats.org/package/2006/content-types'
    types=etree.fromstring(parts['[Content_Types].xml'])
    if not any(e.get('Extension')=='ttf' for e in types):
        etree.SubElement(types,'{'+ns+'}Default',Extension='ttf',ContentType='application/x-font-ttf')
        parts['[Content_Types].xml']=etree.tostring(types,xml_declaration=True,encoding='UTF-8',standalone=True)
        with ZipFile(path,'w',compression=8) as z:
            for n,data in parts.items():z.writestr(n,data)
    doc=Document(path)
    for name in ['First Paragraph', 'Body Text', 'Bibliography', 'Compact']:
        if name not in doc.styles:
            doc.styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH).base_style = doc.styles['normal']
    if 'Caption' not in doc.styles:
        doc.styles.add_style('Caption',WD_STYLE_TYPE.PARAGRAPH).base_style=doc.styles['normal']
    # Pandoc references Hyperlink even when the reference document omits it.
    if 'Hyperlink' not in doc.styles:
        doc.styles.add_style('Hyperlink', WD_STYLE_TYPE.CHARACTER)
    doc.styles['Hyperlink'].font.name = 'Old Standard TT'
    doc.styles['Hyperlink'].font.size = Pt(11)
    # The template uses direct Old Standard TT formatting; its Normal style alone is Arial.
    for name in ['normal','Normal','Body Text','First Paragraph','Compact','Block Text','Caption']:
        if name in doc.styles:
            s=doc.styles[name];s.font.name='Old Standard TT';s.font.color.rgb=RGBColor(0,0,0)
            s.font.size=Pt(9 if name=='Caption' else 11)
    for name in ['Title','Heading 1','Heading 2','Heading 3','Heading 4']:
        if name in doc.styles:doc.styles[name].font.name='Old Standard TT';doc.styles[name].font.color.rgb=RGBColor(0,0,0)
    for i,p in enumerate(doc.paragraphs):
        if i==0 and p.style.name.startswith('Heading'):p.style=doc.styles['Title']
        if p.text in ('References','Appendix'):p.paragraph_format.page_break_before=True
        if p.style.name.startswith('Heading'):p.paragraph_format.keep_with_next=True;p.paragraph_format.space_before=Pt(10)
        if re.match(r'^(?:Figure (?:\d+|\[[^]]+\])|Table \d+|Monitor-design figure)\.',p.text):
            p.style=doc.styles['Caption'];p.paragraph_format.keep_together=True
            p.paragraph_format.line_spacing=1
            p.paragraph_format.space_before=Pt(0)
            p.paragraph_format.space_after=Pt(0)
            # Pandoc can leave captions as normal paragraphs following pictures.
            if i and doc.paragraphs[i-1]._p.xpath('.//w:drawing'):doc.paragraphs[i-1].paragraph_format.keep_with_next=True
        if p._p.xpath('.//w:drawing'):
            p.paragraph_format.space_before=Pt(0)
            p.paragraph_format.space_after=Pt(0)
            p.paragraph_format.line_spacing=1
        p.paragraph_format.widow_control=True
        for r in p.runs:set_run_font(r)
    width=doc.sections[0].page_width.inches-doc.sections[0].left_margin.inches-doc.sections[0].right_margin.inches
    for t in doc.tables:
        n=len(t.columns)
        fractions=([.27,.19,.22,.16,.16] if t.cell(0,0).text=='Scenario' else [.22,.18,.20,.20,.20]) if n==5 else ([.1,.16,.23,.15,.21,.15] if n==6 else ([.10,.15,.75] if n==3 and t.cell(0,0).text=='Round' else [1/n]*n))
        t.autofit=False;set_table_borders(t);set_repeat_header(t.rows[0])
        following = t._tbl.getnext()
        caption_follows = following is not None and ''.join(following.itertext()).startswith('Table ')
        indent=t._tbl.tblPr.find(qn('w:tblInd'))
        if indent is not None:indent.set(qn('w:w'),'0')
        for col,f in zip(t.columns,fractions):col.width=Inches(width*f)
        for j,row in enumerate(t.rows):
            row._tr.get_or_add_trPr().append(OxmlElement('w:cantSplit'))
            for c,f in zip(row.cells,fractions):
                c.width=Inches(width*f);c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
                set_cell_margins(c,top=65,bottom=65,start=70,end=70)
                if j==0:set_cell_shading(c,'E2E8F0')
                elif j%2==0:set_cell_shading(c,'F4F4F4')
                for p in c.paragraphs:
                    p.style=doc.styles['normal'];p.paragraph_format.space_after=Pt(0);p.paragraph_format.line_spacing=1
                    if caption_follows and j == len(t.rows)-1:
                        p.paragraph_format.keep_with_next=True
                    for r in p.runs:
                        set_run_font(r,size=9)
                        if j==0:r.bold=True
    footnote_id = apply_layout(doc, template)
    doc.save(path)
    copy_title_footnote(path, template, footnote_id)


def build(args):
    source=args.source.resolve();out=args.output.resolve();out.parent.mkdir(parents=True,exist_ok=True)
    raw=source.read_bytes();snapshot=out.with_suffix('.source.md')
    # Resolve images against the manuscript directory, even though the saved snapshot is in build/.
    ast,images=prepare_ast(source,raw)
    bib=args.bibliography.resolve()
    if not bib.exists():
        if not args.allow_missing_bibliography:raise FileNotFoundError(f'{bib}; wait for the bibliography, or use --allow-missing-bibliography for a provisional layout only')
        bib=out.parent/'empty.bib';bib.write_text('')
    temporary = tempfile.NamedTemporaryFile(suffix='.docx', dir=out.parent, delete=False)
    temporary.close()
    draft = Path(temporary.name)
    command=['pandoc','--from=json','--to=docx','--standalone','--citeproc','-M','link-citations=true','-M','link-bibliography=true',f'--bibliography={bib}',
             f'--csl={args.csl.resolve()}',f'--reference-doc={args.template.resolve()}',
             f'--resource-path={source.parent}',f'--output={draft}']
    result=subprocess.run(command,input=json.dumps(ast),text=True,capture_output=True,check=True)
    out.with_suffix('.pandoc.log').write_text(result.stderr)
    if 'not found' in result.stderr and not args.allow_missing_bibliography:
        draft.unlink(missing_ok=True)
        raise RuntimeError(result.stderr)
    format_docx(draft, args.template)
    d=Document(draft);all_text='\n'.join(''.join(t.text or '' for t in n.iter(qn('w:t'))) for n in d._element.body)
    unresolved=re.findall(r'\[@[^\]]+\]|\?\?\?',all_text)
    if unresolved and not args.allow_missing_bibliography:raise ValueError(f'Unresolved citations: {unresolved}')
    draft.replace(out)
    snapshot.write_bytes(raw)
    command[-1] = f'--output={out}'
    manifest={'source':str(source.relative_to(ROOT)),'source_sha256':hashlib.sha256(raw).hexdigest(),
              'template_sha256':digest(args.template),'bibliography_sha256':digest(bib),'csl_sha256':digest(args.csl),
              'images':images,'markdown_tables':sum(b.get('t')=='Table' for b in walk(ast)),
              'pandoc_command':command,'pandoc_version':subprocess.check_output(['pandoc','--version'],text=True).splitlines()[0],
              'provisional_missing_bibliography':not args.bibliography.exists(),'unresolved_citation_markers':unresolved,
              'source_changed_during_build':source.read_bytes()!=raw}
    out.with_suffix('.build.json').write_text(json.dumps(manifest,indent=2)+'\n')
    return out


def render(output,renderer):
    qa=output.parent/(output.stem+'-pages')
    if qa.exists():
        for old in qa.glob('page-*.png'):old.unlink()
    if renderer is not None:
        subprocess.run([sys.executable,str(renderer),str(output),'--output_dir',str(qa),'--emit_pdf'],check=True)
    else:
        qa.mkdir(parents=True, exist_ok=True)
        soffice = shutil.which('soffice') or shutil.which('libreoffice')
        if not soffice or not shutil.which('pdftoppm'):
            raise RuntimeError('Install LibreOffice (soffice) and Poppler (pdftoppm) on PATH; see report/BUILD.md')
        rendered_pdf = qa/output.with_suffix('.pdf').name
        rendered_pdf.unlink(missing_ok=True)
        with tempfile.TemporaryDirectory(prefix='paper-lo-') as profile:
            subprocess.run([soffice, f'-env:UserInstallation={Path(profile).as_uri()}',
                            '--headless', '--norestore', '--convert-to', 'pdf',
                            '--outdir', str(qa), str(output)], check=True, timeout=180)
        if not rendered_pdf.is_file():
            raise RuntimeError('LibreOffice did not produce a PDF')
        subprocess.run(['pdftoppm', '-r', '120', '-png', str(rendered_pdf), str(qa/'page')], check=True)
        for png in qa.glob('page-*.png'):
            normalized = qa/f'page-{int(png.stem.split("-")[-1])}.png'
            if png != normalized:
                png.rename(normalized)
    pdf=output.with_suffix('.pdf');shutil.copy2(qa/pdf.name,pdf)
    return measure_pages(output)


def measure_pages(output):
    pdf = output.with_suffix('.pdf')
    pages=[p.extract_text() for p in PdfReader(pdf).pages]
    marker=re.compile(r'^Code and Data\s*$',re.M)
    boundary=next(i for i,p in enumerate(pages) if marker.search(p))
    before=marker.split(pages[boundary])[0].strip()
    body_pages=boundary+bool(before)
    bbox = ET.fromstring(subprocess.check_output(['pdftotext', '-bbox-layout', str(pdf), '-']))
    ns = {'x': 'http://www.w3.org/1999/xhtml'}
    body_words = []
    for page_number, page in enumerate(bbox.findall('.//x:page', ns), 1):
        words = page.findall('.//x:word', ns)
        stop = next((i for i in range(len(words)-2)
                     if [w.text for w in words[i:i+3]] == ['Code', 'and', 'Data']), None)
        body_words.extend((page_number, float(w.get('yMax'))) for w in words[:stop])
        if stop is not None:
            break
    last_page = body_words[-1][0]
    last_y = max(y for page, y in body_words if page == last_page)
    section = Document(output).sections[0]
    top = section.top_margin.pt
    text_height = section.page_height.pt-top-section.bottom_margin.pt
    page_equivalents = last_page-1+(last_y-top)/text_height
    result={'main_body_page_equivalents':round(page_equivalents, 3),
            'main_body_last_text_y_pt':last_y,
            'fraction_definition':'Whole preceding pages plus occupied text-area height on the final body page.',
            'total_pages':len(pages),'main_body_pages':body_pages,'main_body_limit':8,
            'within_limit':body_pages<=8,'code_and_data_starts_page':boundary+1,
            'references_starts_page':next((i+1 for i,p in enumerate(pages) if re.search(r'^References\s*$',p,re.M)),None),
            'count_definition':'Physical pages containing Title through Conclusion, inclusive; partial final page counts.',
            'page_text_file':output.with_suffix('.pages.txt').name}
    output.with_suffix('.pages.txt').write_text('\n\f\n'.join(pages))
    output.with_suffix('.pagination.json').write_text(json.dumps(result,indent=2)+'\n')
    return result


if __name__=='__main__':
    parser=argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--source',type=Path,default=ROOT/'report/paper.md')
    parser.add_argument('--template',type=Path,default=ROOT/'report/template.docx')
    parser.add_argument('--bibliography',type=Path,default=ROOT/'report/references.bib')
    parser.add_argument('--csl',type=Path,default=ROOT/'report/apa.csl')
    parser.add_argument('--output',type=Path,default=ROOT/'report/build/paper.docx')
    parser.add_argument('--renderer',type=Path,default=RENDERER)
    parser.add_argument('--allow-missing-bibliography',action='store_true')
    parser.add_argument('--no-render',action='store_true')
    args=parser.parse_args();out=build(args)
    if not args.no_render:print(json.dumps(render(out,args.renderer),indent=2))
