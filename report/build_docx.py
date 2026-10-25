#!/usr/bin/env python3
"""Build the paper layout test from Markdown and the supplied Word template."""

from __future__ import annotations

import argparse
import re
import shutil
import subprocess
import tempfile
from pathlib import Path

import docx
from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
DEFAULT_TEMPLATE = ROOT / "report" / "template.docx"
DEFAULT_METHODS = ROOT / "report" / "draft_methods.md"
DEFAULT_FIGURE = ROOT / "runs" / "20261025-1850-e03-behaviour-summary" / "cumulative_give_in.png"
DEFAULT_OUTPUT = ROOT / "report" / "build" / "methods_layout_test.docx"

SECTION_ORDER = [
    "1. Introduction",
    "2. Related Work",
    "3. Methods",
    "4. Results",
    "5. Discussion and Limitations",
    "6. Conclusion",
    "Code and Data",
    "Author Contributions (optional)",
    "References",
    "Appendix (optional)",
    "LLM Usage Statement",
]

NEAR_TERM_ROWS = [
    ["Near-term A", "Pooled", "0.586 [0.541, 0.631]", "0.618 [0.571, 0.666]",
     "0.650 [0.608, 0.688]", "0.557 [0.518, 0.593]"],
    ["Near-term A", "LOSO", "0.579 [0.537, 0.619]", "0.579 [0.537, 0.619]",
     "0.650 [0.608, 0.688]", "0.557 [0.518, 0.593]"],
]


def set_cell_shading(cell, color: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), color)


def set_cell_margins(cell, top=80, start=90, bottom=80, end=90) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for edge, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_borders(table, color="D9D9D9", size="8") -> None:
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        node = borders.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            borders.append(node)
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), size)
        node.set(qn("w:color"), color)


def set_repeat_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    header = OxmlElement("w:tblHeader")
    header.set(qn("w:val"), "true")
    tr_pr.append(header)


def set_run_font(run, name="Old Standard TT", size=None, color=None) -> None:
    run.font.name = name
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:hAnsi"), name)
    if size is not None:
        run.font.size = Pt(size)
    if color is not None:
        run.font.color.rgb = RGBColor.from_string(color)


INLINE_PATTERN = re.compile(r"(`[^`]+`|\*\*[^*]+\*\*|\*[^*]+\*)")


def add_inline(paragraph, text: str) -> None:
    text = text.replace("\\$", "$")
    for part in INLINE_PATTERN.split(text):
        if not part:
            continue
        run = paragraph.add_run()
        if part.startswith("**") and part.endswith("**"):
            run.text = part[2:-2]
            run.bold = True
        elif part.startswith("*") and part.endswith("*"):
            run.text = part[1:-1]
            run.italic = True
        elif part.startswith("`") and part.endswith("`"):
            run.text = part[1:-1]
            run.italic = True
        else:
            run.text = part
        set_run_font(run)


def clear_paragraph(paragraph) -> None:
    for child in list(paragraph._p):
        if child.tag != qn("w:pPr"):
            paragraph._p.remove(child)


def prepare_front_matter(doc: Document) -> None:
    for text_node in doc._element.body.xpath(".//w:t"):
        value = text_node.text or ""
        if value == "PROJECT TITLE":
            text_node.text = "Recruited"
        elif value.startswith("Author name ") or value == "Affiliation":
            text_node.text = ""

    abstract_table = doc.tables[0].cell(1, 0).tables[-1]
    paragraph = abstract_table.cell(0, 0).paragraphs[0]
    clear_paragraph(paragraph)
    paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    run = paragraph.add_run(
        "We test whether a target agent's internal activations predict when it will give in to "
        "a peer agent's repeated requests to violate an operator rule. The layout below exercises "
        "the current Methods prose, a representative results table, and a full-width figure."
    )
    run.italic = True
    set_run_font(run)


def clear_body_after_front_matter(doc: Document) -> None:
    body = doc._element.body
    children = list(body.iterchildren())
    for child in children[3:]:
        if child.tag != qn("w:sectPr"):
            body.remove(child)


def parse_methods(path: Path):
    lines = path.read_text(encoding="utf-8").splitlines()
    blocks = []
    paragraph_lines = []

    def flush() -> None:
        if paragraph_lines:
            blocks.append(("paragraph", " ".join(paragraph_lines).strip()))
            paragraph_lines.clear()

    for line in lines:
        stripped = line.strip()
        if not stripped:
            flush()
            continue
        if stripped.startswith("# 3. Methods"):
            continue
        if stripped.startswith("*Draft prose for Section 3"):
            continue
        if stripped.startswith(">"):
            flush()
            continue
        if stripped.startswith("## "):
            flush()
            blocks.append(("heading", stripped[3:]))
        elif stripped.startswith("- "):
            flush()
            blocks.append(("bullet", stripped[2:]))
        else:
            paragraph_lines.append(stripped)
    flush()
    return blocks


def add_methods(doc: Document, path: Path) -> None:
    for kind, text in parse_methods(path):
        if kind == "heading":
            paragraph = doc.add_paragraph(style="Heading 3")
            paragraph.paragraph_format.keep_with_next = True
            add_inline(paragraph, text)
        elif kind == "bullet":
            paragraph = doc.add_paragraph(style="normal")
            paragraph.paragraph_format.left_indent = Inches(0.25)
            paragraph.paragraph_format.first_line_indent = Inches(-0.16)
            paragraph.paragraph_format.space_after = Pt(3)
            add_inline(paragraph, f"• {text}")
        else:
            paragraph = doc.add_paragraph(style="normal")
            paragraph.paragraph_format.space_after = Pt(6)
            paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            add_inline(paragraph, text)


def add_caption(doc: Document, label: str, text: str, before=True):
    paragraph = doc.add_paragraph(style="normal")
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_before = Pt(8)
    paragraph.paragraph_format.space_after = Pt(4 if before else 8)
    paragraph.paragraph_format.keep_with_next = before
    lead = paragraph.add_run(f"{label}. ")
    lead.bold = True
    set_run_font(lead, size=9.5)
    run = paragraph.add_run(text)
    set_run_font(run, size=9.5)
    return paragraph


def add_results_table(doc: Document) -> None:
    add_caption(
        doc,
        "Table 1",
        "Near-term violation prediction on the combined Qwen3-32B dataset. AUROC compares "
        "steps within the same scenario and round; brackets show 95% family-bootstrap confidence intervals.",
        before=True,
    )
    headers = [
        "Label", "Analysis", "Probe within scenario x round", "Fold-stratified probe",
        "Jev forecaster L3", "GPT-5-nano detector L2",
    ]
    widths = [0.72, 0.65, 1.40, 1.18, 1.20, 1.35]
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_table_borders(table)
    set_repeat_header(table.rows[0])
    for index, (cell, header, width) in enumerate(zip(table.rows[0].cells, headers, widths)):
        cell.width = Inches(width)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        set_cell_margins(cell)
        set_cell_shading(cell, "404040")
        paragraph = cell.paragraphs[0]
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragraph.paragraph_format.space_after = Pt(0)
        run = paragraph.add_run(header)
        run.bold = True
        set_run_font(run, size=8.5, color="FFFFFF")
    for row_index, values in enumerate(NEAR_TERM_ROWS):
        row = table.add_row()
        for col_index, (cell, value, width) in enumerate(zip(row.cells, values, widths)):
            cell.width = Inches(width)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell)
            if row_index % 2:
                set_cell_shading(cell, "F2F2F2")
            paragraph = cell.paragraphs[0]
            paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT if col_index < 2 else WD_ALIGN_PARAGRAPH.CENTER
            paragraph.paragraph_format.space_after = Pt(0)
            run = paragraph.add_run(value)
            set_run_font(run, size=8.5)


def add_results_figure(doc: Document, figure_path: Path) -> None:
    intro = doc.add_paragraph(style="normal")
    intro.paragraph_format.space_before = Pt(10)
    intro.paragraph_format.space_after = Pt(4)
    add_inline(
        intro,
        "Figure 1 shows the cumulative give-in rate by round across target models and pressure conditions.",
    )
    paragraph = doc.add_paragraph(style="normal")
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_after = Pt(2)
    paragraph.paragraph_format.keep_with_next = True
    shape = paragraph.add_run().add_picture(str(figure_path), width=Inches(6.25))
    shape._inline.docPr.set(
        "descr",
        "Six line charts comparing cumulative give-in rates by round for Qwen3-32B, Gemma4-31B, and gpt-oss-20b under plain and higher-pressure peer conditions.",
    )
    shape._inline.docPr.set("title", "Cumulative give-in rate by round")
    add_caption(
        doc,
        "Figure 1",
        "Cumulative give-in rate by round for the target-model behavior screens. Curves use the "
        "same deterministic full-value publication rule as the main evaluation.",
        before=False,
    )


def add_section_heading(doc: Document, text: str) -> None:
    paragraph = doc.add_paragraph(style="Heading 2")
    paragraph.paragraph_format.keep_with_next = True
    add_inline(paragraph, text)


def build(template: Path, methods: Path, figure: Path, output: Path) -> None:
    doc = Document(template)
    prepare_front_matter(doc)
    clear_body_after_front_matter(doc)
    doc.add_page_break()

    for section in SECTION_ORDER:
        add_section_heading(doc, section)
        if section == "3. Methods":
            add_methods(doc, methods)
        elif section == "4. Results":
            paragraph = doc.add_paragraph(style="normal")
            paragraph.paragraph_format.space_after = Pt(6)
            add_inline(
                paragraph,
                "Table 1 reports the representative all-data near-term comparison used to test table layout.",
            )
            add_results_table(doc)
            add_results_figure(doc, figure)
        elif section == "5. Discussion and Limitations":
            subheading = doc.add_paragraph(style="Heading 3")
            subheading.paragraph_format.keep_with_next = True
            add_inline(subheading, "Limitations")
            future = doc.add_paragraph(style="Heading 3")
            future.paragraph_format.keep_with_next = True
            add_inline(future, "Future Work")

    for style_name in ("Title", "Heading 2", "Heading 3"):
        style = doc.styles[style_name]
        style.font.color.rgb = RGBColor(0, 0, 0)

    output.parent.mkdir(parents=True, exist_ok=True)
    doc.save(output)


def export_pdf(docx_path: Path, output_dir: Path) -> Path:
    soffice = shutil.which("soffice")
    if not soffice:
        raise RuntimeError("soffice not found; prepend the bundled runtime bin directory to PATH")
    with tempfile.TemporaryDirectory(prefix="apa83-lo-") as profile:
        subprocess.run(
            [
                soffice,
                "--headless",
                f"-env:UserInstallation=file://{profile}",
                "--convert-to",
                "pdf",
                "--outdir",
                str(output_dir),
                str(docx_path),
            ],
            check=True,
        )
    return output_dir / f"{docx_path.stem}.pdf"


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--template", type=Path, default=DEFAULT_TEMPLATE)
    parser.add_argument("--methods", type=Path, default=DEFAULT_METHODS)
    parser.add_argument("--figure", type=Path, default=DEFAULT_FIGURE)
    parser.add_argument("--output", type=Path, default=DEFAULT_OUTPUT)
    parser.add_argument("--pdf", action="store_true")
    args = parser.parse_args()
    build(args.template.resolve(), args.methods.resolve(), args.figure.resolve(), args.output.resolve())
    print(f"Build engine: python-docx {docx.__version__}; Pandoc not used")
    print(args.output.resolve())
    if args.pdf:
        print(export_pdf(args.output.resolve(), args.output.resolve().parent))


if __name__ == "__main__":
    main()
